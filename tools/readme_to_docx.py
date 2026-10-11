"""Turn README.md into a Word document. One-off export, not part of the app."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "README.md"
TARGET = ROOT / "TrentPlayer-README.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(paragraph, text: str, base_size: int = 11) -> None:
    parts = INLINE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(base_size)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(base_size)
        if run.font.name != "Consolas":
            run.font.name = "Calibri"


def shade_header(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): "1F2937",
        },
    )
    tcPr.append(shd)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r_index, row in enumerate(rows):
        for c_index in range(width):
            value = row[c_index] if c_index < len(row) else ""
            cell = table.rows[r_index].cells[c_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_runs(paragraph, value.strip(), 10)
            if r_index == 0:
                shade_header(cell)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    document.add_paragraph()


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for code_line in block:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                run = paragraph.add_run(code_line if code_line else " ")
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(10)
            document.add_paragraph()
            continue

        if stripped.startswith("|"):
            table_rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [cell.strip() for cell in raw.split("|")]
                if not all(set(cell) <= set("-: ") for cell in cells):
                    table_rows.append(cells)
                i += 1
            add_table(document, table_rows)
            continue

        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:].strip())
            i += 1
            continue

        paragraph = document.add_paragraph()
        add_runs(paragraph, stripped)
        i += 1

    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
